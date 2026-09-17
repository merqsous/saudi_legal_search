import io
import re
from datetime import datetime
from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from api.db import query_all, query_one, get_db
from api.routes.cases import _get_user_from_auth, get_case_for_access
from api.embeddings import get_client

router = APIRouter()

MODEL = "gpt-4o-mini"

DOC_TYPES = {
    "pleading": "لائحة دعوى (صحيفة دعوى)",
    "memo": "مذكرة دفاع / مرافعة",
    "letter": "خطاب رسمي",
    "legal_opinion": "رأي قانوني مكتوب",
}

DRAFT_SYSTEM_PROMPT = """أنت "مساعد الباحث" — صائغ مستندات قانونية محترف متخصص في القانون السعودي.
اكتب مستنداً قانونياً بناءً على معلومات القضية والمستندات المرتبطة بها أدناه.

نوع المستند المطلوب: {doc_type}
تعليمات المحامي: {instructions}

معلومات القضية:
{case_context}

{judgments_context}

المتطلبات:
1. اكتب بالعربية الفصحى بأسلوب قانوني رسمي مطابق للممارسة القضائية السعودية.
2. استخدم العناوين والترقيم المناسب لنوع المستند (مقدمة، وقائع، طلبات...).
3. اذكر أطراف القضية وأرقامها كما وردت في معلومات القضية.
4. عند الاستشهاد بالأحكام المرتبطة، اذكر أرقامها ومحاكمها.
5. لا تخترع معلومات غير موجودة في السياق المرفق.
6. ابدأ المستند بعنوانه الرئيسي في السطر الأول.
7. اختم بعبارة "والله ولي التوفيق" إن كان المستند دعوى أو مذكرة مرافعة."""


def init_case_drafts_tables():
    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute("""
                CREATE TABLE IF NOT EXISTS case_drafts (
                    id SERIAL PRIMARY KEY,
                    case_id INTEGER NOT NULL REFERENCES user_cases(id) ON DELETE CASCADE,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    doc_type VARCHAR(30) NOT NULL,
                    instructions TEXT,
                    content TEXT NOT NULL,
                    created_at TIMESTAMP DEFAULT NOW()
                );
            """)
            cur.execute("CREATE INDEX IF NOT EXISTS idx_case_drafts_case_id ON case_drafts(case_id);")
            cur.close()
    except Exception as e:
        import logging
        logging.error(f"init_case_drafts_tables failed: {e}")


class DraftRequest(BaseModel):
    doc_type: str  # pleading | memo | letter | legal_opinion
    instructions: str | None = None


def _get_draft_for_user(draft_id: int, user_id: int) -> dict:
    row = query_one(
        """SELECT d.* FROM case_drafts d
           JOIN user_cases uc ON d.case_id = uc.id
           WHERE d.id = %s AND (
               uc.user_id = %s
               OR (uc.firm_id IS NOT NULL AND uc.firm_id IN (
                   SELECT fm.firm_id FROM firm_members fm WHERE fm.user_id = %s))
           );""",
        [draft_id, user_id, user_id],
    )
    if not row:
        raise HTTPException(status_code=404, detail="المستند غير موجود")
    return dict(row)


@router.get("/cases/{case_id}/drafts")
def list_drafts(case_id: int, authorization: str = Header(None)):
    user_id = _get_user_from_auth(authorization)
    if not user_id:
        raise HTTPException(status_code=401, detail="غير مصرح")
    if not get_case_for_access(case_id, user_id):
        raise HTTPException(status_code=404, detail="القضية غير موجودة")

    rows = query_all(
        """SELECT id, doc_type, instructions, created_at,
                  length(content) AS content_length
           FROM case_drafts WHERE case_id = %s ORDER BY created_at DESC;""",
        [case_id],
    )
    return {"drafts": [dict(r) for r in rows]}


@router.post("/cases/{case_id}/drafts")
def create_draft(case_id: int, req: DraftRequest, authorization: str = Header(None)):
    """Generate a legal document (pleading, memo, letter, opinion) from case context."""
    user_id = _get_user_from_auth(authorization)
    if not user_id:
        raise HTTPException(status_code=401, detail="غير مصرح")
    if not get_case_for_access(case_id, user_id):
        raise HTTPException(status_code=404, detail="القضية غير موجودة")
    if req.doc_type not in DOC_TYPES:
        raise HTTPException(status_code=400, detail="نوع المستند غير صحيح")

    from api.routes.case_chat import _build_case_context
    case_context = _build_case_context(case_id)

    judgments = query_all(
        """SELECT j.judgment_number, j.judgment_year, j.full_text,
                  ct.name_ar AS court_type, cl.name_ar AS court_level, l.city_ar AS city
           FROM case_judgments cj
           JOIN judgments j ON cj.judgment_id = j.id
           LEFT JOIN cases c ON j.case_id = c.id
           LEFT JOIN court_types ct ON c.court_type_id = ct.id
           LEFT JOIN locations l ON c.location_id = l.id
           LEFT JOIN court_levels cl ON j.court_level_id = cl.id
           WHERE cj.case_id = %s ORDER BY cj.added_at DESC LIMIT 5;""",
        [case_id],
    )
    judgments_context = ""
    if judgments:
        judgments_context = "أحكام مرتبطة بالقضية يمكن الاستشهاد بها (مقتطفات):\n"
        for j in judgments:
            excerpt = (j.get("full_text") or "")[:2000]
            judgments_context += (
                f"\n- حكم رقم {j['judgment_number']} — {j.get('court_type') or ''} "
                f"{j.get('court_level') or ''} — {j.get('city') or ''} — {j.get('judgment_year') or ''}:\n"
                f"{excerpt}\n"
            )

    prompt = DRAFT_SYSTEM_PROMPT.format(
        doc_type=DOC_TYPES[req.doc_type],
        instructions=req.instructions or "بدون تعليمات إضافية — صغ المستند وفقاً لمعلومات القضية.",
        case_context=case_context,
        judgments_context=judgments_context,
    )

    try:
        client = get_client()
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.4,
            max_tokens=2500,
        )
        content = resp.choices[0].message.content.strip()
    except Exception as e:
        print(f"[CASE_DRAFT] AI failed: {e}")
        raise HTTPException(
            status_code=503,
            detail="الخدمة الذكية غير متاحة مؤقتاً. يرجى المحاولة لاحقاً.",
        )

    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(
            """INSERT INTO case_drafts (case_id, user_id, doc_type, instructions, content)
               VALUES (%s, %s, %s, %s, %s) RETURNING id, created_at;""",
            (case_id, user_id, req.doc_type, req.instructions, content),
        )
        row = cur.fetchone()
        cur.close()
    return {"status": "ok", "draft_id": row[0], "content": content, "created_at": row[1]}


@router.get("/cases/drafts/{draft_id}")
def get_draft(draft_id: int, authorization: str = Header(None)):
    draft = _get_draft_for_user(draft_id, _get_user_from_auth(authorization))
    return draft


@router.delete("/cases/drafts/{draft_id}")
def delete_draft(draft_id: int, authorization: str = Header(None)):
    user_id = _get_user_from_auth(authorization)
    if not user_id:
        raise HTTPException(status_code=401, detail="غير مصرح")
    _get_draft_for_user(draft_id, user_id)
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM case_drafts WHERE id = %s;", [draft_id])
        cur.close()
    return {"status": "ok"}


@router.get("/cases/drafts/{draft_id}/export/docx")
def export_draft_docx(draft_id: int, authorization: str = Header(None), token: str = None):
    """Export a generated draft as a branded Word document."""
    raw = token or authorization
    if not raw:
        raise HTTPException(status_code=401, detail="غير مصرح")
    raw = raw.replace("Bearer ", "")
    from api.routes.auth import get_session as _get_session
    session = _get_session(raw)
    if not session:
        raise HTTPException(status_code=401, detail="غير مصرح")

    draft = _get_draft_for_user(draft_id, session["user_id"])

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed")

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(12)
    style_pPr = style.element.get_or_add_pPr()
    style_pPr.append(style_pPr.makeelement(qn('w:bidi'), {}))

    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section._sectPr.append(section._sectPr.makeelement(qn('w:bidi'), {}))

    def _set_rtl(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pPr.append(pPr.makeelement(qn('w:bidi'), {}))

    def _add_run(paragraph, text, size=12, color=None, bold=False):
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        run.bold = bold
        rPr = run._r.get_or_add_rPr()
        rPr.append(rPr.makeelement(qn('w:rtl'), {}))
        return run

    GREEN = RGBColor(0x0B, 0x4B, 0x2A)
    MUTED = RGBColor(0x65, 0x75, 0x6B)

    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(hp, "الباحث — Albaheth", size=14, color=GREEN, bold=True)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(date_para)
    _add_run(date_para, f"التاريخ: {datetime.now().strftime('%Y-%m-%d')}", size=10, color=MUTED)
    doc.add_paragraph()

    for line in draft["content"].split("\n"):
        line = line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        _set_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
        clean = re.sub(r'^#{1,6}\s+', '', clean)
        is_heading = bool(re.match(r'^#{1,6}\s+', line)) or bool(re.match(r'^\d+\.\s+\S', line))
        _add_run(p, clean, size=13 if is_heading else 12,
                 color=GREEN if is_heading else None, bold=is_heading)

    doc.add_paragraph()
    p = doc.add_paragraph()
    _set_rtl(p)
    _add_run(p, "هذا المستند تم توليده بواسطة مساعد الباحث الذكي — يخضع لمراجعة المحامي قبل الاعتماد.",
             size=9, color=MUTED)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="draft_{draft_id}.docx"'},
    )
